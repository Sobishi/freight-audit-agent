from openpyxl import load_workbook
from data.loader import WORKBOOK_PATH


def main():
    print("EXCEL AUDIT FORMULA INSPECTION")
    print("------------------------------")

    workbook = load_workbook(
        WORKBOOK_PATH,
        data_only=False,
        read_only=True,
    )

    audit_sheet = workbook["Audit"]

    # The worksheet headers are on Excel row 4.
    header_row = 4

    # Inspect formulas for the first two invoice lines.
    invoice_rows = [5, 6]

    for row_number in invoice_rows:
        print(f"\nEXCEL ROW {row_number}")
        print("------------------------------")

        for column_number in range(1, 49):
            header = audit_sheet.cell(
                row=header_row,
                column=column_number,
            ).value

            cell = audit_sheet.cell(
                row=row_number,
                column=column_number,
            )

            value = cell.value

            if cell.data_type == "f":
                print(
                    f"{cell.coordinate} | "
                    f"{header} | "
                    f"{value}"
                )

    workbook.close()


if __name__ == "__main__":
    main()
