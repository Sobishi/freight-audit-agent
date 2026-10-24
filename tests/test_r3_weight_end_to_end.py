import json
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def test_cli_generates_weight_validation_report_and_run_metadata(tmp_path):
    config_dir = tmp_path / "config"
    carriers_dir = config_dir / "carriers"
    mappings_dir = config_dir / "mappings"

    carriers_dir.mkdir(parents=True)
    mappings_dir.mkdir(parents=True)

    (carriers_dir / "swiftline.yaml").write_text(
        (PROJECT_ROOT / "config/carriers/swiftline.yaml").read_text()
    )

    (mappings_dir / "swiftline_invoice.yaml").write_text(
        """
carrier_id: swiftline
carrier_name: Swiftline
header_row: 0
date_format: "%d/%m/%Y"
decimal: "."
thousands: ","
columns:
  "Inv Ref": Invoice Ref
  "Waybill No": AWB
  "Dispatch Date": Ship Date
  "Product": Service Type
  "Zone": Billed Zone
  "Chargeable Kg": Billed Weight
  "Billed Unit": Billed Weight Unit
  "Actual Kg": Actual Weight
  "Actual Unit": Weight Unit
  "Length Cm": Length
  "Width Cm": Width
  "Height Cm": Height
  "Dimension Unit": Dimension Unit
  "Package Count": Pieces
  "Freight": Base Freight
  "Surcharge": Surcharge Amount
  "VAT": Tax Amount
  "Amount Due": Total
""".strip()
        + "\n"
    )

    invoice_path = tmp_path / "invoice.csv"

    pd.DataFrame([{
        "Inv Ref": "INV001",
        "Waybill No": "AWB001",
        "Dispatch Date": "24/10/2026",
        "Product": "Express",
        "Zone": "A",
        "Chargeable Kg": 5.0,
        "Billed Unit": "kg",
        "Actual Kg": 3.0,
        "Actual Unit": "kg",
        "Length Cm": 40.0,
        "Width Cm": 30.0,
        "Height Cm": 20.0,
        "Dimension Unit": "cm",
        "Package Count": 1,
        "Freight": 100.0,
        "Surcharge": 0.0,
        "VAT": 0.0,
        "Amount Due": 100.0,
    }]).to_csv(invoice_path, index=False)

    report_path = tmp_path / "audit_report.xlsx"
    log_dir = tmp_path / "logs"

    command = [
        sys.executable,
        str(PROJECT_ROOT / "main.py"),
        str(invoice_path),
        "--carrier",
        "swiftline",
        "--config-dir",
        str(config_dir),
        "--weight-tolerance",
        "0.1",
        "--package-count-verified",
        "--output",
        str(report_path),
        "--log-dir",
        str(log_dir),
    ]

    result = subprocess.run(
        command,
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, (
        result.stdout + "\n" + result.stderr
    )

    assert report_path.exists()

    workbook = load_workbook(report_path, read_only=True)

    try:
        assert "Invoice Data" in workbook.sheetnames
        assert "Audit Summary" in workbook.sheetnames
        assert "Weight Validation" in workbook.sheetnames

        sheet = workbook["Weight Validation"]
        rows = list(sheet.values)

        headers = rows[0]
        values = rows[1]
        record = dict(zip(headers, values))

        assert record["AWB"] == "AWB001"
        assert record["Weight Status"] == "calculated"
        assert record["Expected Chargeable Weight (kg)"] == 4.0
        assert record["Billed Weight (kg)"] == 5.0
        assert record["Weight Difference (kg)"] == 1.0
        assert record["Weight Comparison Status"] == "billed_weight_high"

        assert "Potential Exposure" not in headers
        assert "Combined Potential Exposure" not in headers

    finally:
        workbook.close()

    log_files = list(log_dir.glob("*.json"))

    assert len(log_files) == 1

    run_data = json.loads(log_files[0].read_text())

    assert run_data["settings"]["weight_validation_enabled"] is True
    assert run_data["settings"]["weight_tolerance_kg"] == 0.1
    assert run_data["settings"]["package_count_verified"] is True

    assert (
        run_data["summary"]["weight_status_counts"]["calculated"]
        == 1
    )


def test_cli_without_weight_options_preserves_existing_report(tmp_path):
    invoice_path = tmp_path / "invoice.csv"

    pd.DataFrame([{
        "Invoice Ref": "INV002",
        "AWB": "AWB002",
        "Ship Date": "2026-10-24",
        "Service Type": "Express",
        "Billed Weight": 5.0,
        "Base Freight": 100.0,
        "Total": 100.0,
    }]).to_csv(invoice_path, index=False)

    report_path = tmp_path / "ordinary_audit.xlsx"

    result = subprocess.run(
        [
            sys.executable,
            str(PROJECT_ROOT / "main.py"),
            str(invoice_path),
            "--output",
            str(report_path),
            "--log-dir",
            str(tmp_path / "logs"),
        ],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, (
        result.stdout + "\n" + result.stderr
    )

    workbook = load_workbook(report_path, read_only=True)

    try:
        assert "Audit Summary" in workbook.sheetnames
        assert "Audit Findings" in workbook.sheetnames
        assert "Invoice Data" in workbook.sheetnames
        assert "Weight Validation" not in workbook.sheetnames
    finally:
        workbook.close()


def test_incomplete_measurements_produce_review_not_failure(tmp_path):
    from audit_engine.weight_reporting import (
        generate_weight_validation_results,
    )
    from data.carrier_config import load_carrier_config

    config = load_carrier_config("swiftline")

    invoice = pd.DataFrame([{
        "Invoice Line ID": "LINE001",
        "AWB": "AWB003",
        "Carrier ID": "swiftline",
        "Actual Weight": 3.0,
        "Weight Unit": "kg",
        "Length": 40.0,
        "Width": None,
        "Height": 20.0,
        "Dimension Unit": "cm",
        "Pieces": 1,
        "Billed Weight": 5.0,
        "Billed Weight Unit": "kg",
    }])

    results = generate_weight_validation_results(
        invoice,
        profile=config.profile,
        package_count_verified=True,
        weight_tolerance=0.1,
    )

    assert len(results) == 1
    assert results.iloc[0]["Weight Status"] == "incomplete_dimensions"
    assert not bool(results.iloc[0]["Weight Calculation Eligible"])
    assert pd.isna(
        results.iloc[0]["Expected Chargeable Weight (kg)"]
    )
    assert pd.isna(
        results.iloc[0]["Weight Comparison Status"]
    )
