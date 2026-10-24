import pandas as pd
from openpyxl import load_workbook

from reporting.excel_report import generate_excel_report


def test_weight_evidence_is_reported_without_changing_exposure(tmp_path):
    invoice = pd.DataFrame({
        "Invoice Line ID": ["line-001"],
        "AWB": ["AWB001"],
        "Total": [120.0],
    })

    findings = pd.DataFrame(columns=[
        "Invoice Line ID",
        "Finding Type",
        "Potential Exposure",
    ])

    summary = {
        "Total Invoice Lines": 1,
        "Gross Potential Exposure": 0.0,
    }

    weight_results = pd.DataFrame({
        "Invoice Line ID": ["line-001"],
        "Weight Status": ["calculation_review"],
        "Weight Review Reason": ["missing_billed_weight_unit"],
        "Expected Chargeable Weight (kg)": [5.0],
        "Billed Weight": [7.0],
        "Billed Weight Unit": [pd.NA],
    })

    output = tmp_path / "r3_weight_report.xlsx"

    generate_excel_report(
        invoice,
        findings,
        summary,
        str(output),
        weight_results=weight_results,
    )

    workbook = load_workbook(output, data_only=True)

    assert "Weight Validation" in workbook.sheetnames

    sheet = workbook["Weight Validation"]
    rows = list(sheet.values)

    headers = rows[0]
    values = dict(zip(headers, rows[1]))

    assert values["Invoice Line ID"] == "line-001"
    assert values["Weight Status"] == "calculation_review"
    assert values["Weight Review Reason"] == "missing_billed_weight_unit"
    assert values["Expected Chargeable Weight (kg)"] == 5.0

    summary_sheet = workbook["Audit Summary"]
    summary_values = {
        row[0]: row[1]
        for row in summary_sheet.iter_rows(
            min_row=2,
            min_col=1,
            max_col=2,
            values_only=True,
        )
        if row[0] is not None
    }

    assert summary_values["Gross Potential Exposure"] == 0.0
