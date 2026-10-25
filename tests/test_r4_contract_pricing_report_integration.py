import pandas as pd
from openpyxl import load_workbook

from reporting.excel_report import generate_excel_report


def test_contract_pricing_report_preserves_existing_exposure(tmp_path):
    invoice = pd.DataFrame({
        "Invoice Line ID": ["line-001"],
        "AWB": ["AWB001"],
        "Total": [120.0],
    })

    findings = pd.DataFrame(columns=[
        "Invoice Line ID",
        "Finding Type",
        "Potential Exposure",
    ])

    summary = {
        "Total Invoice Lines": 1,
        "Gross Potential Exposure": 0.0,
    }

    contract_results = pd.DataFrame([{
        "Invoice Line ID": "line-001",
        "AWB": "AWB001",
        "Carrier ID": "swiftline",
        "Pricing Status": "potential_overcharge",
        "Pricing Review Reason": None,
        "Expected Chargeable Weight (kg)": 5.0,
        "Expected Base Freight": 100.0,
        "Invoiced Base Freight": 120.0,
        "Potential Contract Overcharge": 20.0,
        "Rate Difference": 20.0,
        "Currency": "USD",
    }])

    output = tmp_path / "r4_contract_pricing.xlsx"

    generate_excel_report(
        invoice,
        findings,
        summary,
        str(output),
        contract_pricing_results=contract_results,
    )

    workbook = load_workbook(output, data_only=True)

    assert "Contract Pricing" in workbook.sheetnames

    rows = list(workbook["Contract Pricing"].values)
    assert len(rows) == 2

    values = dict(zip(rows[0], rows[1]))

    assert values["Invoice Line ID"] == "line-001"
    assert values["Pricing Status"] == "potential_overcharge"
    assert values["Expected Chargeable Weight (kg)"] == 5.0
    assert values["Expected Base Freight"] == 100.0
    assert values["Invoiced Base Freight"] == 120.0
    assert values["Potential Contract Overcharge"] == 20.0
    assert values["Rate Difference"] == 20.0
    assert values["Currency"] == "USD"

    summary_values = {
        row[0]: row[1]
        for row in workbook["Audit Summary"].iter_rows(
            min_row=2,
            min_col=1,
            max_col=2,
            values_only=True,
        )
        if row[0] is not None
    }

    assert summary_values["Gross Potential Exposure"] == 0.0
    assert "Combined Potential Exposure" not in summary_values


def test_contract_pricing_sheet_is_optional(tmp_path):
    output = tmp_path / "without_r4.xlsx"

    generate_excel_report(
        pd.DataFrame({
            "Invoice Line ID": ["line-001"],
            "AWB": ["AWB001"],
        }),
        pd.DataFrame(columns=["Finding Type"]),
        {"Total Invoice Lines": 1},
        str(output),
    )

    workbook = load_workbook(output, data_only=True)

    assert "Contract Pricing" not in workbook.sheetnames
