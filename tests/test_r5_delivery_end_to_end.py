"""R5 delivery evidence end-to-end integration."""

import json
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def test_cli_delivery_evidence_preserves_duplicate_awb_lines(tmp_path):
    invoice_path = tmp_path / "invoice.csv"
    evidence_path = tmp_path / "evidence.xlsx"
    report_path = tmp_path / "delivery_audit.xlsx"
    log_dir = tmp_path / "logs"

    pd.DataFrame([
        {
            "AWB": "NBX700",
            "Carrier": "Northbridge Express",
            "Total": 100.0,
        },
        {
            "AWB": "NBX700",
            "Carrier": "Northbridge Express",
            "Total": 120.0,
        },
    ]).to_csv(invoice_path, index=False)

    with pd.ExcelWriter(evidence_path) as writer:
        pd.DataFrame([
            {
                "AWB": "NBX700",
                "Carrier": "Northbridge Express",
                "Status": "Delivered",
                "Delivery Status": "Delivered",
                "Delivery Date": "2026-10-20",
                "Signed By": "Reception",
            }
        ]).to_excel(
            writer,
            sheet_name="Carrier Tracking",
            index=False,
        )

    result = subprocess.run(
        [
            sys.executable,
            str(PROJECT_ROOT / "main.py"),
            str(invoice_path),
            "--delivery-evidence",
            str(evidence_path),
            "--output",
            str(report_path),
            "--log-dir",
            str(log_dir),
        ],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, (
        result.stdout + "\n" + result.stderr
    )

    assert "DELIVERY EVIDENCE" in result.stdout
    assert report_path.exists()

    workbook = load_workbook(
        report_path,
        read_only=True,
        data_only=True,
    )

    try:
        assert "Delivery Evidence" in workbook.sheetnames

        rows = list(workbook["Delivery Evidence"].values)

        records = [
            dict(zip(rows[0], row))
            for row in rows[1:]
        ]

        assert len(records) == 2

        assert [row["Invoice Line ID"] for row in records] == [1, 2]

        assert all(
            row["AWB"] == "NBX700"
            for row in records
        )

        assert all(
            row["Delivery Status"] == "confirmed"
            for row in records
        )

        assert all(
            row["Matched Evidence Count"] == 1
            for row in records
        )

        assert "Contract Pricing" not in workbook.sheetnames
        assert "Weight Validation" not in workbook.sheetnames

    finally:
        workbook.close()

    log_files = list(log_dir.glob("*.json"))

    assert len(log_files) == 1

    run_data = json.loads(log_files[0].read_text())

    assert run_data["status"] == "succeeded"

    assert run_data["summary"]["delivery_status_counts"] == {
        "confirmed": 2
    }

    evidence_inputs = [
        item
        for item in run_data["inputs"]
        if item["role"] == "delivery_evidence"
    ]

    assert len(evidence_inputs) == 1
    assert evidence_inputs[0].get("sha256")


def test_cli_accepts_fourth_row_delivery_evidence_headers(tmp_path):
    invoice_path = tmp_path / "invoice.csv"
    evidence_path = tmp_path / "evidence.xlsx"
    report_path = tmp_path / "delivery_report.xlsx"

    pd.DataFrame([
        {
            "AWB": "NBX800",
            "Carrier": "Northbridge Express",
            "Total": 100.0,
        }
    ]).to_csv(invoice_path, index=False)

    with pd.ExcelWriter(evidence_path) as writer:
        pd.DataFrame([
            ["Carrier Tracking Extract", None, None],
            ["Tracking results", None, None],
            [None, None, None],
            ["AWB", "Status", "Carrier"],
            ["NBX800", "Delivered", "Northbridge Express"],
        ]).to_excel(
            writer,
            sheet_name="Carrier Tracking",
            header=False,
            index=False,
        )

    result = subprocess.run(
        [
            sys.executable,
            str(PROJECT_ROOT / "main.py"),
            str(invoice_path),
            "--delivery-evidence",
            str(evidence_path),
            "--output",
            str(report_path),
            "--log-dir",
            str(tmp_path / "logs"),
        ],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, (
        result.stdout + "\n" + result.stderr
    )

    workbook = load_workbook(report_path, read_only=True)

    try:
        rows = list(workbook["Delivery Evidence"].values)
        record = dict(zip(rows[0], rows[1]))

        assert record["AWB"] == "NBX800"
        assert record["Delivery Status"] == "confirmed"
        assert record["Matched Evidence Count"] == 1

    finally:
        workbook.close()
