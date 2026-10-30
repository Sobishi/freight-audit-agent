"""R5 delivery evidence Excel report integration tests."""

import pandas as pd
from openpyxl import load_workbook

from reporting.excel_report import generate_excel_report


def make_report_inputs():
    invoice = pd.DataFrame([
        {
            "Invoice Line ID": 1,
            "AWB": "NBX500",
            "Carrier": "Northbridge Express",
            "Total": 150.0,
        }
    ])

    findings = pd.DataFrame([
        {
            "Invoice Line ID": 1,
            "AWB": "NBX500",
            "Finding": "Test finding",
        }
    ])

    summary = {
        "Total Invoice Lines": 1,
        "Combined Potential Exposure": 25.0,
    }

    delivery_results = pd.DataFrame([
        {
            "Invoice Line ID": 1,
            "AWB": "NBX500",
            "Delivery Status": "review",
            "Matched Evidence Count": 0,
            "Evidence Sources": "",
            "Has Signature Name": False,
            "Review Reason": "carrier_mismatch",
        }
    ])

    return invoice, findings, summary, delivery_results


def test_excel_report_includes_delivery_evidence_when_supplied(tmp_path):
    invoice, findings, summary, delivery_results = make_report_inputs()

    output = tmp_path / "delivery_report.xlsx"

    generate_excel_report(
        audited_invoice=invoice,
        findings=findings,
        summary=summary,
        output_path=str(output),
        delivery_evidence_results=delivery_results,
    )

    workbook = load_workbook(output, read_only=True)

    try:
        assert "Delivery Evidence" in workbook.sheetnames

        rows = list(workbook["Delivery Evidence"].values)

        assert rows[0] == tuple(delivery_results.columns)
        assert rows[1][0] == 1
        assert rows[1][1] == "NBX500"
        assert rows[1][2] == "review"
        assert rows[1][6] == "carrier_mismatch"
    finally:
        workbook.close()


def test_excel_report_omits_delivery_sheet_when_not_supplied(tmp_path):
    invoice, findings, summary, _ = make_report_inputs()

    output = tmp_path / "standard_report.xlsx"

    generate_excel_report(
        audited_invoice=invoice,
        findings=findings,
        summary=summary,
        output_path=str(output),
    )

    workbook = load_workbook(output, read_only=True)

    try:
        assert "Delivery Evidence" not in workbook.sheetnames
    finally:
        workbook.close()


def test_delivery_evidence_does_not_change_financial_summary(tmp_path):
    invoice, findings, summary, delivery_results = make_report_inputs()

    output = tmp_path / "exposure_report.xlsx"

    generate_excel_report(
        audited_invoice=invoice,
        findings=findings,
        summary=summary,
        output_path=str(output),
        delivery_evidence_results=delivery_results,
    )

    workbook = load_workbook(output, read_only=True)

    try:
        rows = list(workbook["Audit Summary"].values)

        metrics = {
            row[0]: row[1]
            for row in rows[1:]
            if len(row) >= 2 and row[0] is not None
        }

        assert metrics["Combined Potential Exposure"] == 25.0
        assert "Delivery Evidence" in workbook.sheetnames
    finally:
        workbook.close()


def test_delivery_evidence_preserves_existing_financial_metrics(tmp_path):
    invoice, findings, summary, delivery_results = make_report_inputs()

    standard_path = tmp_path / "standard_financial_report.xlsx"
    delivery_path = tmp_path / "delivery_financial_report.xlsx"

    generate_excel_report(
        audited_invoice=invoice,
        findings=findings,
        summary=summary,
        output_path=str(standard_path),
    )

    generate_excel_report(
        audited_invoice=invoice,
        findings=findings,
        summary=summary,
        output_path=str(delivery_path),
        delivery_evidence_results=delivery_results,
    )

    def read_summary(path):
        workbook = load_workbook(path, read_only=True)

        try:
            rows = list(workbook["Audit Summary"].values)

            return {
                row[0]: row[1]
                for row in rows[1:]
                if len(row) >= 2 and row[0] is not None
            }
        finally:
            workbook.close()

    standard_summary = read_summary(standard_path)
    delivery_summary = read_summary(delivery_path)

    assert delivery_summary == standard_summary
    assert delivery_summary["Combined Potential Exposure"] == 25.0
